from pathlib import Path
import json,time,math
import win32com.client
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE,MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

INK='10162F';CREAM='F8F5EE';CYAN='24C3D9';CORAL='FF7058';LIME='C5F66A';VIOLET='7958F5';GRAY='B4B8CB';DARKGRAY='4C5064'
W,H=13.333333,7.5
def rgb(c):return RGBColor.from_string(c)
class Deck:
 def __init__(self):
  self.p=Presentation();self.p.slide_width=Inches(W);self.p.slide_height=Inches(H);self.pages=[];self.s=None;self.events=[];self.color=INK;self.bg=CREAM
 def page(self,scene,duration,bg=CREAM,title='',speech='',tag='',source=''):
  self.s=self.p.slides.add_slide(self.p.slide_layouts[6]);self.s.background.fill.solid();self.s.background.fill.fore_color.rgb=rgb(bg);self.bg=bg;self.color=CREAM if bg in [INK,VIOLET,DARKGRAY] else INK
  self.events=[];self.pages.append({'scene':scene,'seconds':duration,'title':title,'speech':speech,'events':self.events,'bg':bg,'tag':tag,'source':source})
  self.text(tag,.55,.28,11,.3,12,self.color,font='Consolas',name='section-tag')
  self.text(f'{scene:02}',12.0,.28,.75,.3,13,self.color,font='Consolas',align=PP_ALIGN.RIGHT)
  if source:self.text(source,.55,7.09,12.0,.23,10,self.color,name='source')
  return self.s
 def text(self,t,x,y,w,h,size=24,color=None,bold=False,font='Aptos',name=None,align=None):
  h=max(h,(size*1.18*(t.count('\n')+1)+2)/72)
  q=self.s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));q.name=name or f'text-{len(self.s.shapes)}';tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
  for i,line in enumerate(t.split('\n')):
   p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color or self.color);p.space_after=Pt(0)
   if align is not None:p.alignment=align
  return q.name
 def rect(self,x,y,w,h,fill,name=None,line=None,kind=MSO_SHAPE.RECTANGLE,opacity=1):
  q=self.s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.name=name or f'shape-{len(self.s.shapes)}'
  if fill:q.fill.solid();q.fill.fore_color.rgb=rgb(fill)
  else:q.fill.background()
  if line:q.line.color.rgb=rgb(line);q.line.width=Pt(1.5)
  else:q.line.fill.background()
  return q.name
 def dot(self,x,y,r=.08,color=CYAN,name=None):return self.rect(x-r,y-r,2*r,2*r,color,name,kind=MSO_SHAPE.OVAL)
 def ring(self,x,y,r,color=CYAN,name=None):return self.rect(x-r,y-r,2*r,2*r,None,name,line=color,kind=MSO_SHAPE.OVAL)
 def line(self,x,y,xx,yy,color=None,width=1.5,arrow=False,name=None):
  q=self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(xx),Inches(yy));q.name=name or f'line-{len(self.s.shapes)}';q.line.color.rgb=rgb(color or self.color);q.line.width=Pt(width)
  if arrow:
   e=OxmlElement('a:tailEnd');e.set('type','triangle');q.line._get_or_add_ln().append(e)
  return q.name
 def anim(self,name,kind='fade',at=0,duration=.6,**kwargs):
  self.events.append(dict(shape=name,kind=kind,at=at,duration=duration,**kwargs));return name
 def move(self,name,dx,dy,at,duration,repeat=1):return self.anim(name,'motion',at,duration,dx=dx,dy=dy,repeat=repeat)
 def pulse(self,name,at=0,duration=1.4,repeat=1,scale=112):return self.anim(name,'scale',at,duration,scale=scale,repeat=repeat,reverse=True)
 def packet(self,x,y,xx,yy,at=0,duration=2,repeat=1,color=CYAN,name=None):
  n=self.dot(x,y,.055,color,name);self.anim(n,'fade',at,.12);self.move(n,xx-x,yy-y,at+.12,duration,repeat);self.anim(n,'exit',at+.12+duration*repeat,.12);return n
 def title(self,t,x=.55,y=.9,w=12,h=1.8,size=50,color=None,name='hero-title'):
  return self.text(t,x,y,w,h,size,color,True,'Arial',name)
 def save(self,path):
  for i,s in enumerate(self.p.slides):s.notes_slide.notes_text_frame.text=f"Cảnh{self.pages[i]['scene']} |{self.pages[i]['seconds']}giây\n"+self.pages[i]['speech']+'\nNguồn:'+self.pages[i]['source']
  self.p.save(path);Path(path).with_suffix('.motion.json').write_text(json.dumps(self.pages,ensure_ascii=False,indent=2),encoding='utf-8')

def native_render(pptx,main_count=None,video=None,render_dir=None,pdf=None):
 pptx=Path(pptx).resolve();pages=json.loads(pptx.with_suffix('.motion.json').read_text(encoding='utf-8'))
 a=win32com.client.DispatchEx('PowerPoint.Application');p=None;stats=[]
 try:
  p=a.Presentations.Open(str(pptx),False,False,False)
  for i,page in enumerate(pages,1):
   s=p.Slides(i);seq=s.TimeLine.MainSequence
   while seq.Count:seq(1).Delete()
   for ev in sorted(page['events'],key=lambda v:v['at']):
    q=s.Shapes(ev['shape']);kind=ev['kind']
    if kind in ['motion','scale']:
     e=seq.AddEffect(q,0,0,2)
     if kind=='motion':b=e.Behaviors.Add(1);b.MotionEffect.Path=f"M 0 0 L {ev['dx']/W:.8f} {ev['dy']/H:.8f} E"
     else:
      b=e.Behaviors.Add(3);b.ScaleEffect.FromX=100;b.ScaleEffect.FromY=100;b.ScaleEffect.ToX=ev.get('scale',110);b.ScaleEffect.ToY=ev.get('scale',110)
    else:
     e=seq.AddEffect(q,{'fade':10,'exit':10,'wipe':22,'appear':1,'fly':2}[kind],0,2)
     if kind=='exit':e.Exit=True
     if kind=='wipe':e.EffectParameters.Direction=4
     if kind=='fly':e.EffectParameters.Direction=4
    e.Timing.Duration=ev['duration'];e.Timing.TriggerDelayTime=ev['at']
    if ev.get('repeat',1)>1:e.Timing.RepeatCount=ev['repeat']
    if ev.get('reverse'):e.Timing.AutoReverse=True
    try:e.Timing.Accelerate=.15;e.Timing.Decelerate=.15
    except:pass
   s.SlideShowTransition.EntryEffect=0;s.SlideShowTransition.AdvanceOnTime=True;s.SlideShowTransition.AdvanceOnClick=True;s.SlideShowTransition.AdvanceTime=page['seconds']
   if main_count and i>main_count:s.SlideShowTransition.Hidden=True;s.SlideShowTransition.AdvanceOnTime=False
   stats.append({'slide':i,'scene':page['scene'],'effects':seq.Count,'seconds':page['seconds']})
  if main_count:p.SlideShowSettings.RangeType=2;p.SlideShowSettings.StartingSlide=1;p.SlideShowSettings.EndingSlide=main_count;p.SlideShowSettings.AdvanceMode=2
  p.Save()
  if render_dir:
   r=Path(render_dir).resolve();r.mkdir(exist_ok=True,parents=True)
   for i in range(1,p.Slides.Count+1):p.Slides(i).Export(str(r/f'slide_{i:02}.png'),'PNG',1600,900)
  if pdf:
   hidden=[p.Slides(i).SlideShowTransition.Hidden for i in range(1,p.Slides.Count+1)]
   for i in range(1,p.Slides.Count+1):p.Slides(i).SlideShowTransition.Hidden=False
   p.SlideShowSettings.RangeType=1
   p.SaveAs(str(Path(pdf).resolve()),32)
   for i,v in enumerate(hidden,1):p.Slides(i).SlideShowTransition.Hidden=v
  if video:
   if main_count:
    for i in range(p.Slides.Count,main_count,-1):p.Slides(i).Delete()
   p.SlideShowSettings.RangeType=1
   p.CreateVideo(str(Path(video).resolve()),True,5,1080,30,90);start=time.time();last=0
   while p.CreateVideoStatus in [1,2]:
    time.sleep(1)
    if time.time()-last>30:print('EXPORT',Path(video).name,round(time.time()-start),flush=True);last=time.time()
   assert p.CreateVideoStatus==3,p.CreateVideoStatus
  pptx.with_suffix('.native_stats.json').write_text(json.dumps(stats,indent=2),encoding='utf-8')
 finally:
  if p is not None:p.Saved=True;p.Close()
  a.Quit()
