from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
import win32com.client,time,json
O=Path(__file__).resolve().parent;D=O/'concepts';D.mkdir(exist_ok=True)
p=Presentation();p.slide_width=Inches(13.333);p.slide_height=Inches(7.5);s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string('10162F')
q=s.shapes.add_shape(9,Inches(1),Inches(3),Inches(.3),Inches(.3));q.name='packet';q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string('24C3D9')
t=s.shapes.add_textbox(Inches(1),Inches(1),Inches(11),Inches(1));t.name='hero';t.text='MOTION / 118';t.text_frame.paragraphs[0].font.size=Pt(80);t.text_frame.paragraphs[0].font.color.rgb=RGBColor.from_string('F8F5EE')
p.save(D/'motion_probe.pptx')
a=win32com.client.DispatchEx('PowerPoint.Application');pr=None
try:
 pr=a.Presentations.Open(str(D/'motion_probe.pptx'),False,False,False);s=pr.Slides(1);seq=s.TimeLine.MainSequence
 e=seq.AddEffect(s.Shapes('hero'),10,0,2);e.Timing.Duration=.6;e.Timing.TriggerDelayTime=0
 e=seq.AddEffect(s.Shapes('packet'),0,0,2);b=e.Behaviors.Add(1);b.MotionEffect.Path='M 0 0 L 0.7 0 E';e.Timing.Duration=2;e.Timing.TriggerDelayTime=1;e.Timing.RepeatCount=3
 s.SlideShowTransition.AdvanceOnTime=True;s.SlideShowTransition.AdvanceTime=8;s.SlideShowTransition.AdvanceOnClick=False
 pr.Save();pr.CreateVideo(str(D/'motion_probe.mp4'),True,8,720,30,80)
 while pr.CreateVideoStatus in [1,2]:time.sleep(1)
 print('PROBE_STATUS',pr.CreateVideoStatus,flush=True)
finally:
 if pr is not None:pr.Saved=True;pr.Close()
 a.Quit()
