from pathlib import Path
import json, csv, math, hashlib, html
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from content import MAIN, QA, SCENES

OUT=Path(__file__).resolve().parent; ROOT=OUT.parents[1]
P=Presentation(); P.slide_width=Inches(13.333333); P.slide_height=Inches(7.5)
NAVY='0D3560'; BLUE='1769AA'; CYAN='139DCB'; INK='172B43'; MUTED='526A81'; PALE='EFF6FC'; LIGHT='D6E7F5'; ORANGE='B65B29'; WHITE='FFFFFF'
def rgb(h):return RGBColor.from_string(h)
def shape(s,x,y,w,h,fill=PALE,line=None,name=None,kind=MSO_SHAPE.RECTANGLE):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h)); q.fill.solid();q.fill.fore_color.rgb=rgb(fill)
 if line:q.line.color.rgb=rgb(line);q.line.width=Pt(1.2)
 else:q.line.fill.background()
 if name:q.name=name
 return q
def text(s,t,x,y,w,h,size=20,color=INK,bold=False,align=None,name=None):
 # Keep the full Aptos line box inside the editable text frame. Explicit
 # paragraph breaks also carry 5pt spacing; top anchoring preserves pixels.
 explicit_lines=t.count('\n')+1
 h=max(h,(1.2*size*explicit_lines+5*(explicit_lines-1)+1)/72)
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));q.name=name or 'Text '+str(len(s.shapes))
 tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color)
  p.space_after=Pt(5)
  if align is not None:p.alignment=align
 return q
def line(s,x1,y1,x2,y2,color=LIGHT,width=2,arrow=False,name=None):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));q.line.color.rgb=rgb(color);q.line.width=Pt(width)
 if arrow:
  e=OxmlElement('a:tailEnd');e.set('type','triangle');q.line._get_or_add_ln().append(e)
 if name:q.name=name
 return q
def node(s,t,x,y,w=1.7,h=.75,active=True,name=None):
 q=shape(s,x,y,w,h,BLUE if active else PALE,name=name)
 text(s,t,x+.12,y+.15,w-.24,h-.18,18,WHITE if active else MUTED,True,PP_ALIGN.CENTER,name=(name+'-text') if name else None)
 return q
def footer(s,scene,src,idx):
 text(s,src,.58,7.03,11.6,.24,10,MUTED)
 text(s,f'{idx:02}',12.1,7.0,.6,.26,12,MUTED,align=PP_ALIGN.RIGHT)
 # A subject-relevant flow rail carries scene progress, no title underline.
 line(s,.6,6.78,12.72,6.78,LIGHT,1,name='!!scene-flow-rail')
 for k in range(8):shape(s,.6+k*1.7,6.735,.09,.09,BLUE if k+1<=scene else LIGHT,name=f'!!scene-stage-{k}',kind=MSO_SHAPE.OVAL)
 shape(s,.56+(scene-1)*1.7,6.70,.16,.16,CYAN,name='!!scene-flow',kind=MSO_SHAPE.OVAL)
def base(title,scene,src,idx):
 s=P.slides.add_slide(P.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(WHITE)
 text(s,title,.6,.55,12.1,1.2,36,NAVY,True,name=f'#s{idx}-title')
 footer(s,scene,src,idx);return s
def note(s,txt):s.notes_slide.notes_text_frame.text=txt
def tree(s,x,y,scale=1,name='tree'):
 for j,(dx,dy) in enumerate([(0,0),(-.4,.45),(.4,.45),(-.62,.9),(-.18,.9),(.2,.9),(.62,.9)]):
  if j:line(s,x+(0 if j<3 else (-.4 if j<5 else .4))*scale,y+(0 if j<3 else .45)*scale,x+dx*scale,y+dy*scale,BLUE,1.5)
  shape(s,x+dx*scale-.06,y+dy*scale-.06,.12,.12,BLUE,name=f'{name}-{j}',kind=MSO_SHAPE.OVAL)

HIST=ROOT/'results/final_evaluation/unsw_nb15/phase9_holdout/tables/final_metrics.csv'
rows=list(csv.DictReader(HIST.open(encoding='utf-8')))
R={r['pipeline']:r for r in rows if r['eval_view']=='official_test'}
AGG=ROOT/'results/controlled_extensions/p21_train_only_science_summary_20261019_022655/aggregate.json'
AG=json.loads(AGG.read_text())
ARM=['A0_RF42_NONE','A1_RF17_NONE','A2_RF42_BALANCED','A3_RF42_OVERSAMPLE','A4_STACK42_NONE']
means={a:{k:sum(v['metrics_by_arm'][a][k] for v in AG['per_seed'].values())/5 for k in ['macro_f1','mcc','fpr','fnr']} for a in ARM}
assert int(R['random_forest_42']['fp'])-int(R['stacking_42']['fp'])==142
assert int(R['stacking_42']['fn'])-int(R['random_forest_42']['fn'])==118
assert sum(d for _,d,_,_,_ in MAIN)==360

def architecture(s,stage):
 # All four keyframes use the same named actors; cyan flow traverses actual stages.
 xs=[.65,3.1,5.55,8.0,10.55]; ys=[3.0]*5
 labels=['Bản ghi\nlưu lượng','Đặc trưng\n42 / 17 biến','LR · DT · RF','Kết hợp\nVoting / Stack','Normal\nhoặc Attack']
 for j in range(4):line(s,xs[j]+2.02,3.47,xs[j+1]-.05,3.47,BLUE if j<stage else LIGHT,2,True,name=f'!!actor-pipe-edge-{j}')
 for j in range(5):
  active=j<=stage or stage==3
  a=shape(s,xs[j],3.0,2.02,1.0,BLUE if active else PALE,name=f'!!actor-pipe-{j}')
  text(s,labels[j],xs[j]+.08,3.17,1.86,.75,18,WHITE if active else MUTED,True,PP_ALIGN.CENTER,name=f'!!actor-pipe-label-{j}')
 # Native vector representation of one illustrative flow, never a measured packet trace.
 fx=[.8,3.7,6.1,11.1][stage]
 shape(s,fx,2.39,.24,.24,CYAN,name='!!actor-flow-record',kind=MSO_SHAPE.OVAL)
 line(s,fx+.12,2.63,fx+.12,2.94,CYAN,2,True,name='!!actor-flow-pointer')
 if stage==0:
  text(s,'Dữ liệu flow đã được trích xuất trong UNSW-NB15',.7,4.65,11.9,.6,24,NAVY,True)
  text(s,'Minh họa một bản ghi • Pipeline offline, chưa phải IDS bắt gói trực tiếp',.7,5.45,11.9,.5,18,MUTED)
 elif stage==1:
  text(s,'Loại id, label, attack_cat khỏi predictor',.7,4.65,11.9,.55,24,NAVY,True)
  text(s,'One-hot: proto / service / state   •   Nhánh LR có chuẩn hóa',.7,5.45,11.9,.5,20,MUTED)
 elif stage==2:
  text(s,'Preprocessing được fit trong phần huấn luyện',.7,4.65,11.9,.55,24,NAVY,True)
  text(s,'Nhánh 17 biến lịch sử có selection bias trước CV; P21 chọn lại trong FIT.',.7,5.4,11.9,.65,18,MUTED)
 else:
  text(s,'Ngưỡng 0,5 đã khóa   →   Attack là lớp dương',.7,4.65,11.9,.55,24,NAVY,True)
  text(s,'FP: Normal → Attack                  FN: Attack → Normal',.7,5.45,11.9,.5,22,MUTED)

def model_flow(s,mode):
 text(s,'Cùng đầu vào',.7,3.1,2,.5,22,NAVY,True)
 for j,l in enumerate(['Logistic\nRegression','Decision\nTree','Random\nForest']):
  yy=2.15+j*1.1;node(s,l,3.3,yy,2.35,.85,name=f'!!actor-learner-{j}');line(s,2.4,3.4,3.2,yy+.4,BLUE,1.6,True)
  line(s,5.75,yy+.45,7.0,3.48,BLUE,1.6,True)
 if mode=='voting':
  node(s,'Luật gộp\ncố định',7.0,3.0,2.05,1,name='!!actor-combiner');text(s,'Hard: đa số nhãn\nSoft: trung bình xác suất',9.55,2.85,3.0,1.55,22,NAVY,True)
  text(s,'Diversity về cách học chưa bảo đảm lợi ích khi gộp',.7,5.75,11.8,.55,22,MUTED)
 else:
  node(s,'LR tầng\nkết hợp',7.0,3.0,2.05,1,name='!!actor-combiner');line(s,9.15,3.48,10.0,3.48,BLUE,2,True);node(s,'Normal /\nAttack',10.1,3.0,2.45,1,name='!!actor-decision')
  text(s,'HUẤN LUYỆN META',.7,1.95,2.5,.3,13,BLUE,True)
  text(s,'Đầu vào meta = xác suất OOF; không dùng dự đoán in-sample',.7,5.5,11.8,.6,22,NAVY,True)
  text(s,'Khi suy luận: base learner đã fit lại → xác suất → meta LR',.7,6.12,11.8,.3,16,MUTED)

def test_bars(s,phase):
 fp=[int(R[a]['fp']) for a in ['random_forest_42','stacking_42']];fn=[int(R[a]['fn']) for a in ['random_forest_42','stacking_42']]
 for group,values,maxv,x,w,col,active in [('fp',fp,11000,.85,5.2,BLUE,True),('fn',fn,1000,7.25,4.9,ORANGE,phase>=2)]:
  text(s,'CẢNH BÁO GIẢ · FP' if group=='fp' else 'BỎ SÓT · FN',x,2.0,w,.4,19,col,True)
  for j,(label,val) in enumerate(zip(['RF42','Stacking42'],values)):
   yy=2.9+j*1.15;text(s,label,x,yy-.3,w*.5,.36,18,MUTED)
   shape(s,x,yy+.15,w*(val/maxv) if active else .015,.35,col if active else LIGHT,name=f'!!actor-{group}-bar-{j}')
   text(s,f'{val:,}'.replace(',','.') if active else '…',x+w*.5,yy-.36,w*.5,.45,23,col if active else MUTED,True,PP_ALIGN.RIGHT,name=f'!!actor-{group}-value-{j}')
  line(s,x,5.15,x+w,5.15,LIGHT,1);text(s,'0',x,5.25,.5,.25,12,MUTED);text(s,f'{maxv:,}'.replace(',','.'),x+w-.8,5.25,.8,.25,12,MUTED,align=PP_ALIGN.RIGHT)
 if phase==1:text(s,'−142 FP',.85,5.85,5,.6,34,BLUE,True)
 else:
  text(s,'−142 FP',.85,5.85,5,.6,34,BLUE,True);text(s,'+118 FN',7.25,5.85,5,.6,34,ORANGE,True)
 text(s,'Hai trục dùng thang riêng; số liệu là số bản ghi trên cùng official TEST.',.85,6.47,11.8,.22,12,MUTED)

plan=[]
for idx,(sc,dur,view,title,spoken) in enumerate(MAIN,1):
 src='Nguồn: báo cáo V6, chương 3–5.'
 if sc==5:src='Nguồn: P9 final_metrics.csv · official TEST, 82.332 dòng · ngưỡng 0,5.'
 if sc in [6,7]:src='Nguồn: P21 aggregate.json · hậu nghiệm trên TRAIN; cùng validation qua 5 FIT seed.'
 if view=='cost':src='Nguồn: V6 §5.1 / P9 confusion counts · phân tích chi phí giả định cho đúng 2 mô hình.'
 s=base(title,sc,src,idx)
 if view=='opening':
  text(s,'NGHIÊN CỨU ENSEMBLE LEARNING\nCHO PHÁT HIỆN XÂM NHẬP MẠNG',.65,2.2,8.5,.8,22,BLUE,True)
  text(s,'Cái giá của từng loại sai lầm',.65,3.55,8.5,.65,30,NAVY,True)
  # A flow divides into two possible outcomes, signature of the talk.
  line(s,9.2,3.85,10.25,3.85,BLUE,3,True);line(s,10.3,3.85,11.5,2.7,BLUE,3,True);line(s,10.3,3.85,11.5,5.0,ORANGE,3,True)
  shape(s,9.0,3.72,.26,.26,CYAN,name='!!actor-hook-flow',kind=MSO_SHAPE.OVAL)
  text(s,'FP',11.6,2.42,1,.5,32,BLUE,True);text(s,'FN',11.6,4.82,1,.5,32,ORANGE,True)
  text(s,'Dương Ngọc Phú · Phạm Thế Vỹ · Vũ Văn Kiên · Lê Hoàng Nguyên Quân',.65,5.65,11.9,.4,16,MUTED)
  text(s,'GVHD: TS. Lê Ngọc Hiếu   |   PTIT cơ sở TP.HCM   |   2026',.65,6.12,11.9,.3,15,MUTED)
 elif view=='question':
  text(s,'Bản ghi\nlưu lượng',.85,2.9,3.2,1.05,28,BLUE,True)
  node(s,'Mô hình IDS',4.6,3.05,3.1,1.05,name='!!actor-ids');line(s,3.35,3.55,4.5,3.55,BLUE,2,True)
  line(s,7.8,3.55,9.1,2.65,BLUE,2,True);line(s,7.8,3.55,9.1,4.2,BLUE,2,True)
  text(s,'Normal',9.4,2.35,3,.65,32,BLUE,True);text(s,'Attack',9.4,3.9,3,.65,32,ORANGE,True)
  text(s,'Tốt hơn về chất lượng dự đoán,\nhay chỉ chuyển sai số?',1,4.8,11.3,1.15,30,NAVY,True,PP_ALIGN.CENTER)
 elif view=='dataset':
  text(s,'175.341',.85,2.25,5.6,1,64,BLUE,True);text(s,'dòng official TRAIN',.9,3.3,5.4,.5,23,NAVY)
  text(s,'82.332',7.05,2.25,5.6,1,64,NAVY,True);text(s,'dòng official TEST',7.1,3.3,5.4,.5,23,NAVY)
  shape(s,.85,4.55,5.5,1.35,PALE);shape(s,7.05,4.55,5.5,1.35,PALE)
  text(s,'P21 · nội bộ TRAIN\nĐối chứng bổ sung, hậu nghiệm',1.05,4.78,5.1,.95,21,BLUE,True)
  text(s,'Kết quả lịch sử\nĐánh giá candidate đã khóa',7.25,4.78,5.1,.95,21,NAVY,True)
  text(s,'Không so trực tiếp các mức điểm giữa hai tập đánh giá.',.9,6.1,11.6,.35,18,MUTED)
 elif view.startswith('architecture'):architecture(s,int(view[-1]))
 elif view=='rf':
  for j,x in enumerate([2.0,4.0,6.0,8.0,10.0]):tree(s,x,2.45,1.1,f'!!actor-tree{j}');line(s,x,3.7,6.0,4.6,BLUE,1.5,True)
  node(s,'Tổng hợp nhiều cây',4.35,4.7,3.3,.8)
  text(s,'Mốc RF42 lịch sử: 300 cây  •  42 predictor gốc',.8,5.85,11.8,.5,24,NAVY,True)
 elif view in ['voting','stacking']:model_flow(s,view)
 elif view=='testscore':
  for j,(key,label) in enumerate([('random_forest_42','RF42'),('stacking_42','Stacking42')]):
   x=.85+j*6.3;text(s,label,x,2.1,5.4,.5,25,MUTED,True)
   text(s,f"{float(R[key]['macro_f1']):.6f}",x,2.9,5.6,1.0,57,BLUE if j==0 else NAVY,True)
   text(s,f"MCC  {float(R[key]['mcc']):.6f}",x,4.2,5.5,.5,27,NAVY)
  text(s,'Macro-F1: Stacking − RF = +0,000497',.85,5.5,11.8,.6,29,NAVY,True)
  text(s,'MCC đi theo hướng ngược lại; không kết luận vượt trội.',.85,6.15,11.8,.35,20,MUTED)
 elif view in ['testfp','testfn']:test_bars(s,1 if view=='testfp' else 2)
 elif view=='testsummary':
  text(s,'−142 FP',.85,2.3,5.5,1,58,BLUE,True);text(s,'+118 FN',7.0,2.3,5.5,1,58,ORANGE,True)
  text(s,'CI 95% của Δ Macro-F1 / Δ MCC đều chứa 0',.85,4.15,11.7,.75,30,NAVY,True)
  text(s,'Chưa thấy ưu thế rõ theo phép đo này.\nKhông đồng nghĩa hai mô hình tương đương.',.85,5.25,11.7,1.0,24,MUTED)
 elif view=='p21design':
  text(s,'5 FIT seed',.85,2.2,4.8,.8,43,BLUE,True);text(s,'1 validation',7.1,2.2,5.5,.8,43,NAVY,True)
  for j in range(5):
   shape(s,.95+j*.75,3.5,.4,.4,BLUE,name=f'!!actor-fit-seed-{j}',kind=MSO_SHAPE.OVAL);line(s,1.15+j*.75,4.0,7.0,4.0,LIGHT,1)
  node(s,'35.069 dòng cố định',7.1,3.5,4.8,.95)
  text(s,'60.000 FIT rows / seed',.9,4.6,5.5,.5,23,BLUE,True)
  text(s,'RF42 · RF17 · Balanced · Oversample · Stacking',.9,5.5,11.7,.65,24,NAVY,True)
  text(s,'Group theo vector trùng; không phải 5 TEST độc lập.',.9,6.14,11.7,.4,20,MUTED)
 elif view=='p21score':
  # Paired seed chart: truncation declared; all actual per-seed values shown as points.
  lo,hi=.9440,.9460;xl,xr=3.35,11.85
  for v in [.9440,.9450,.9460]:
   xx=xl+(v-lo)/(hi-lo)*(xr-xl);line(s,xx,2.2,xx,5.65,LIGHT,1);text(s,f'{v:.3f}',xx-.3,5.83,.8,.3,14,MUTED)
  for j,seed in enumerate(AG['seeds']):
   yy=2.35+j*.66;text(s,str(seed),.85,yy-.1,1.5,.4,18,MUTED)
   vals=[AG['per_seed'][str(seed)]['metrics_by_arm'][a]['macro_f1'] for a in [ARM[0],ARM[4]]]
   pts=[xl+(v-lo)/(hi-lo)*(xr-xl) for v in vals];line(s,min(pts),yy,max(pts),yy,LIGHT,3)
   for xx,col in zip(pts,[BLUE,ORANGE]):shape(s,xx-.065,yy-.065,.13,.13,col,kind=MSO_SHAPE.OVAL)
  text(s,'RF42-Lite  0,945220',.85,1.78,5.8,.4,21,BLUE,True);text(s,'Stacking-Lite  0,944731',7.0,1.78,5.8,.4,21,ORANGE,True)
  text(s,'Macro-F1 · trục rút gọn 0,944–0,946; đường nối cùng FIT seed',.85,6.35,11.9,.25,14,MUTED)
 elif view=='p21cost':
  xx=3.6;ww=7.6;maxv=350
  for yy,label,parts in [(2.65,'RF42-Lite',[(241.119,BLUE)]),(4.1,'RF17-Lite',[(225.250,LIGHT),(84.600,BLUE)])]:
   text(s,label,.85,yy+.05,2.6,.5,24,NAVY,True);start=xx
   for k,(val,col) in enumerate(parts):shape(s,start,yy,ww*val/maxv,.6,col,name=f'!!actor-cost-{label}-{k}');start+=ww*val/maxv
  text(s,'241,119 s',9.05,3.35,3.2,.4,24,BLUE,True)
  text(s,'225,250 s chọn biến + 84,600 s fit = 309,850 s',.85,5.25,11.9,.6,26,NAVY,True)
  text(s,'Thời gian các lời gọi đã đo; chưa phải chi phí end-to-end. Macro-F1 RF17: 0,936751.',.85,6.1,11.9,.4,17,MUTED)
  line(s,xx,4.95,xx+ww,4.95,LIGHT,1);text(s,'0 s',xx,5.02,1,.2,11,MUTED);text(s,'350 s',xx+ww-.8,5.02,1,.2,11,MUTED)
 elif view=='tradeoff':
  # Rates from the same validation only; no historical TEST on this coordinate system.
  x0,y0,w,h=2.35,5.6,8.9,3.1
  line(s,x0,2.15,x0,y0,BLUE,1.6);line(s,x0,y0,11.65,y0,BLUE,1.6)
  for x in [0,4,8,12]:text(s,str(x),x0+x/12*w-.1,y0+.15,.5,.3,14,MUTED)
  for y in [0,2,4]:text(s,str(y),x0-.5,y0-y/5*h-.1,.4,.3,14,MUTED)
  text(s,'FNR (%)',.85,1.95,1.8,.4,18,NAVY,True);text(s,'FPR (%)',10.7,6.12,1.8,.4,18,NAVY,True)
  for arm,label,col,dx,dy in [(ARM[0],'RF42',BLUE,.15,.05),(ARM[2],'Balanced',ORANGE,-1.25,-.65),(ARM[3],'Oversample',CYAN,.15,-.45)]:
   xx=x0+means[arm]['fpr']*100/12*w;yy=y0-means[arm]['fnr']*100/5*h
   shape(s,xx-.08,yy-.08,.16,.16,col,kind=MSO_SHAPE.OVAL);text(s,label,xx+dx,yy+dy,2.1,.4,20,col,True)
  text(s,'Trung bình 5 FIT seed · Normal là lớp ít mẫu trong TRAIN',.85,6.42,11.7,.25,14,MUTED)
 elif view=='cost':
  text(s,'L = FP + r × FN',.85,2.1,11.8,.8,43,BLUE,True)
  text(s,'ΔL(Stacking − RF42) = −142 + 118r',.85,3.25,11.8,.8,34,NAVY,True)
  text(s,'r > 1,2034',.85,4.55,5.6,.9,48,ORANGE,True)
  text(s,'RF42 có tổn thất thấp hơn\ntrong đúng cặp so sánh này',7.0,4.6,5.6,1.05,25,NAVY,True)
  text(s,'r = chi phí FN / chi phí FP · giả định minh họa, chưa đo chi phí SOC thực.',.85,6.14,11.8,.45,18,MUTED)
 elif view=='closing':
  text(s,'Đối chiếu có truy vết',.85,2.25,11.7,.55,30,BLUE,True)
  text(s,'Đối chứng matched · Chi phí chọn biến · Đánh đổi FP/FN',.85,3.05,11.7,.8,24,NAVY)
  text(s,'Giới hạn: selection bias lịch sử;\nchưa validation ngoài, chưa triển khai IDS trực tiếp.',.85,4.25,11.7,1.1,25,MUTED)
  text(s,'Mô hình phù hợp phải kiểm soát được cái giá của sai lầm.',.85,5.9,11.7,.55,25,NAVY,True)
 # All main notes carry exact narration and source scope.
 note(s,f'CẢNH {sc} · KEYFRAME {idx} · {dur} giây (bao gồm chuyển vào)\n{spoken}\n\nAnimation: tự chạy theo timing; bấm tay vẫn được. Dừng 1 giây cuối keyframe; không đọc số thập phân dài.\n{src}\nTruy vết chi tiết: Scientific_Evidence_Map.md')
 plan.append({'slide':idx,'scene':sc,'seconds':dur,'view':view,'title':title,'speech':spoken,'source':src})

for j,(title,q,a,src) in enumerate(QA,1):
 s=base(title,8,'Phụ lục vấn đáp · '+src,len(MAIN)+j)
 text(s,q,.85,2.1,11.6,.85,27,BLUE,True)
 text(s,a,.85,3.25,11.6,2.7,25,NAVY)
 text(s,'Trả lời cốt lõi trong 15–25 giây; chỉ mở khi hội đồng hỏi.',.85,6.2,11.6,.35,16,MUTED)
 s._element.set('show','0');note(s,q+'\n'+a+'\nNguồn: '+src)

# Closing slide navigation links to hidden appendix slides, available during Q&A.
closing=P.slides[len(MAIN)-1]
for j in range(12):
 q=text(closing,f'A{j+1}',.85+j*.91,6.5,.75,.25,11,MUTED)
 q.click_action.target_slide=P.slides[len(MAIN)+j]
for s in list(P.slides)[len(MAIN):]:
 q=text(s,'Về kết luận',10.7,.12,2.0,.25,12,BLUE);q.click_action.target_slide=closing

P.core_properties.title='Nghiên cứu Ensemble Learning cho phát hiện xâm nhập mạng'
P.core_properties.subject='Bảo vệ NCKH · 8 cảnh / 360 giây · UNSW-NB15 · V6'
P.core_properties.author='Nhóm NCKH Ensemble IDS'
P.save(OUT/'NCKH_Ensemble_IDS_Final.pptx')
(OUT/'slide_plan.json').write_text(json.dumps(plan,ensure_ascii=False,indent=2),encoding='utf-8')
# SVG mirrors, independently editable, for every technical keyframe and appendix.
# This is an asset export of native shapes, not the image source of the PPTX.
(OUT/'svg').mkdir(exist_ok=True)
for i,s in enumerate(P.slides,1):
 svg=['<svg xmlns="http://www.w3.org/2000/svg" width="1600" height="900" viewBox="0 0 960 540">','<rect width="960" height="540" fill="white"/>']
 for q in s.shapes:
  x,y,w,h=[v/914400*72 for v in [q.left,q.top,q.width,q.height]]
  if q.has_text_frame:
   for k,p in enumerate(q.text_frame.paragraphs):
    if not p.text:continue
    color=str(p.font.color.rgb) if p.font.color.type else INK;size=p.font.size.pt if p.font.size else 20
    svg.append(f'<text x="{x}" y="{y+size+k*(size+6)}" font-family="Aptos,Arial,sans-serif" font-size="{size}" font-weight="{700 if p.font.bold else 400}" fill="#{color}">{html.escape(p.text)}</text>')
  else:
   try:col=str(q.fill.fore_color.rgb);svg.append(f'<rect x="{x}" y="{y}" width="{w}" height="{h}" fill="#{col}"/>')
   except Exception:svg.append(f'<line x1="{x}" y1="{y}" x2="{x+w}" y2="{y+h}" stroke="#{BLUE}"/>')
 svg.append('</svg>');(OUT/'svg'/f'slide_{i:02}.svg').write_text('\n'.join(svg),encoding='utf-8')
print(json.dumps({'slides':len(P.slides),'main':len(MAIN),'appendix':len(QA),'seconds':sum(x['seconds'] for x in plan),'native_shapes':sum(len(s.shapes) for s in P.slides)}))
